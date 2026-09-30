"""
Generate js/data/f01RecommendationSpec.js from the Recommendation Source of Truth:
    docs/specs/ReMotion復健資料庫_含F01推薦GoalMapping.xlsx

Reads (verbatim, no interpretation):
  - F01_Goal候選池        -> F01_GOALS (goal -> direct / supporting exercise ids)
  - F01_推薦GoalMapping   -> F01_EXERCISE_MAPPING (per-exercise direct / supporting goals, role)
  - F01-01 .. F01-20      -> F01_EXERCISE_SPEC (precautions, evidence_status, training_goal)
  - 推薦規則_MVP          -> F01_RECOMMENDATION_RULES (rule table, for traceability)

Also prints a consistency report between the goal pool and the per-exercise mapping.

Usage (from the repo root):  python scripts/generateF01RecommendationSpec.py
Requires: openpyxl
"""
import hashlib
import io
import json
import os
import sys

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC_REL = "docs/specs/ReMotion復健資料庫_含F01推薦GoalMapping.xlsx"
OUT_REL = "js/data/f01RecommendationSpec.js"


def cell(v):
    return "" if v is None else str(v).strip()


def split_ids(v):
    return [s.strip() for s in cell(v).split("|") if s.strip()]


def main():
    src = os.path.join(ROOT, SRC_REL)
    sha256 = hashlib.sha256(open(src, "rb").read()).hexdigest()
    wb = openpyxl.load_workbook(src, data_only=True)

    goals = []
    for row in list(wb["F01_Goal候選池"].iter_rows(values_only=True))[1:]:
        if not row or not cell(row[0]):
            continue
        goals.append({
            "goalId": cell(row[0]),
            "label": cell(row[1]),
            "systemDefinition": cell(row[2]),
            "directExerciseIds": split_ids(row[3]),
            "supportingExerciseIds": split_ids(row[4]),
            "notes": cell(row[5]),
        })

    # Read F01_推薦GoalMapping by HEADER name (V2 schema: primary_goal_id /
    # direct_goal_ids / supporting_goal_ids — one exercise may be Direct for
    # several goals). Fails loudly if a required column is missing.
    map_rows = list(wb["F01_推薦GoalMapping"].iter_rows(values_only=True))
    headers = [cell(h) for h in map_rows[0]]
    required = ["exercise_id", "exercise_name", "primary_goal_id", "direct_goal_ids", "supporting_goal_ids",
                "recommendation_role", "mapping_basis", "mapping_status"]
    missing_cols = [h for h in required if h not in headers]
    if missing_cols:
        raise SystemExit(f"F01_推薦GoalMapping is missing columns: {missing_cols} (found: {headers})")
    col = {h: headers.index(h) for h in required}
    mapping = {}
    for row in map_rows[1:]:
        if not row or not cell(row[col["exercise_id"]]):
            continue
        mapping[cell(row[col["exercise_id"]])] = {
            "name": cell(row[col["exercise_name"]]),
            "primaryGoalId": cell(row[col["primary_goal_id"]]),
            "directGoalIds": split_ids(row[col["direct_goal_ids"]]),
            "supportingGoalIds": split_ids(row[col["supporting_goal_ids"]]),
            "recommendationRole": cell(row[col["recommendation_role"]]),
            "mappingBasis": cell(row[col["mapping_basis"]]),
            "mappingStatus": cell(row[col["mapping_status"]]),
        }

    spec = {}
    for i in range(1, 21):
        sid = f"F01-{i:02d}"
        fields = {}
        for row in wb[sid].iter_rows(values_only=True):
            if row and cell(row[0]) and row[1] is not None and cell(row[0]) not in fields:
                fields[cell(row[0])] = cell(row[1])
        spec[sid] = {
            "precautions": fields.get("precautions", ""),
            "evidenceStatus": fields.get("evidence_status", ""),
            "trainingGoal": fields.get("training_goal", ""),
        }

    rules = []
    for row in list(wb["推薦規則_MVP"].iter_rows(values_only=True))[1:]:
        if not row or not cell(row[0]):
            continue
        rules.append({"ruleId": cell(row[0]), "stage": cell(row[1]), "input": cell(row[2]), "rule": cell(row[3]),
                      "action": cell(row[4]), "uiReason": cell(row[5]), "ruleType": cell(row[6])})

    # Consistency report (README_推薦: F01_Goal候選池 is the Direct / Supporting
    # source of truth; F01_推薦GoalMapping is its reverse view and must match
    # exactly). Checked in both directions, plus primary_goal_id ∈ direct_goal_ids.
    issues = []
    pool = {g["goalId"]: g for g in goals}
    for g in goals:
        for eid in g["directExerciseIds"]:
            m = mapping.get(eid)
            if not m or g["goalId"] not in m["directGoalIds"]:
                issues.append(f"{g['goalId']} pool lists {eid} as DIRECT; mapping direct_goal_ids = {m['directGoalIds'] if m else 'missing'}")
        for eid in g["supportingExerciseIds"]:
            m = mapping.get(eid)
            if not m or g["goalId"] not in m["supportingGoalIds"]:
                issues.append(f"{g['goalId']} pool lists {eid} as SUPPORTING; mapping supporting_goal_ids = {m['supportingGoalIds'] if m else 'missing'}")
    for eid, m in mapping.items():
        for gid in m["directGoalIds"]:
            if gid not in pool or eid not in pool[gid]["directExerciseIds"]:
                issues.append(f"mapping lists {eid} DIRECT for {gid}; pool {gid} direct_exercise_ids does not")
        for gid in m["supportingGoalIds"]:
            if gid not in pool or eid not in pool[gid]["supportingExerciseIds"]:
                issues.append(f"mapping lists {eid} SUPPORTING for {gid}; pool {gid} supporting_exercise_ids does not")
        if m["primaryGoalId"] not in m["directGoalIds"]:
            issues.append(f"{eid} primary_goal_id {m['primaryGoalId']} is not one of its direct_goal_ids {m['directGoalIds']}")

    header = (
        "/**\n"
        " * GENERATED FILE — do not edit by hand.\n"
        f" * Source of Truth: {SRC_REL}\n"
        f" * Source SHA-256:  {sha256}\n"
        " * Generator: scripts/generateF01RecommendationSpec.py\n"
        " *\n"
        " * F01_推薦GoalMapping / F01_Goal候選池 are ReMotion SYSTEM RECOMMENDATION RULES,\n"
        " * not clinical validation standards (README_推薦). Direct / Supporting are\n"
        " * recommendation match levels. Precautions / evidence status are copied\n"
        " * verbatim from each exercise sheet.\n"
        " */\n"
    )
    js = header
    js += f"export const F01_RECOMMENDATION_SPEC_SOURCE = Object.freeze({json.dumps({'file': SRC_REL, 'sha256': sha256}, ensure_ascii=False)});\n\n"
    js += "/** Goal pool (sheet F01_Goal候選池) — the candidate source used by the engine. */\n"
    js += f"export const F01_GOALS = Object.freeze({json.dumps(goals, ensure_ascii=False, indent=2)});\n\n"
    js += "/** Per-exercise mapping (sheet F01_推薦GoalMapping, reverse view of the pool): primaryGoalId / directGoalIds / supportingGoalIds. */\n"
    js += f"export const F01_EXERCISE_MAPPING = Object.freeze({json.dumps(mapping, ensure_ascii=False, indent=2)});\n\n"
    js += "/** Verbatim fields from sheets F01-01 .. F01-20. */\n"
    js += f"export const F01_EXERCISE_SPEC = Object.freeze({json.dumps(spec, ensure_ascii=False, indent=2)});\n\n"
    js += "/** Rule table (sheet 推薦規則_MVP), kept for traceability. */\n"
    js += f"export const F01_RECOMMENDATION_RULES = Object.freeze({json.dumps(rules, ensure_ascii=False, indent=2)});\n\n"
    js += "/** Differences between F01_Goal候選池 and F01_推薦GoalMapping found at generation time (must be empty; the engine follows F01_Goal候選池). */\n"
    js += f"export const F01_SPEC_CONSISTENCY_ISSUES = Object.freeze({json.dumps(issues, ensure_ascii=False, indent=2)});\n"

    out = os.path.join(ROOT, OUT_REL)
    with io.open(out, "w", encoding="utf-8", newline="\n") as f:
        f.write(js)

    out_stream = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
    out_stream.write(f"wrote {OUT_REL}: {len(goals)} goals, {len(mapping)} mapped exercises, {len(spec)} exercise sheets, {len(rules)} rules\n")
    out_stream.write(f"consistency issues ({len(issues)}):\n" + "".join(f"  - {i}\n" for i in issues))
    out_stream.flush()


if __name__ == "__main__":
    main()
