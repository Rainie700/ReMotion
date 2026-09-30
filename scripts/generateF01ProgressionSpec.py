"""
Generate js/data/f01ProgressionSpec.js from the Phase D5 Source of Truth:
    docs/specs/ReMotion_D5_DecisionRules.xlsx   (D5-0 V1)

Reads (verbatim, no interpretation beyond the documented code mapping):
  - 5xSTS_ChangeEvidence      -> selected MDC reference (referenceMdcMs, MDC% metadata)
  - TrainingExposure          -> none / partial / full day ranges
  - D5_DecisionTable          -> R001..R011 (decision mapped to a code, auto_apply_allowed)
  - ExerciseTransition        -> every transition row WITH its status (Confirmed / Draft / Pending)
  - NextCycleSelectionRules   -> NC-M01 / NC-P01 / NC-A01 / NC-A02 / NC-N01 reason templates

Fails loudly (non-zero exit) when the workbook is not the locked V1: a TBD
decision, a rule that is not Confirmed, the old auto_decision_allowed column,
a missing sheet, or an unknown decision label.

Usage (from the repo root):  python scripts/generateF01ProgressionSpec.py
Requires: openpyxl
"""
import hashlib
import io
import json
import os
import sys

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC_REL = "docs/specs/ReMotion_D5_DecisionRules.xlsx"
OUT_REL = "js/data/f01ProgressionSpec.js"

REQUIRED_SHEETS = [
    "README", "5xSTS_ChangeEvidence", "TrainingExposure", "D5_DecisionTable", "ExerciseTransition",
    "TransitionCoverage", "NextCycleSelectionRules", "D5_RuntimeAndHandoff", "ConsistencyCheck",
]

# The only mapping this generator adds: the sheet's Chinese decision label -> program code.
DECISION_CODE = {"不自動判定": "no_auto_decision", "維持": "maintain", "進階": "progress", "調整": "adjust"}
TRANSITION_TYPES = {"progression", "regression", "same_goal_replacement"}


def cell(v):
    return "" if v is None else str(v).strip()


def fail(msg):
    raise SystemExit(f"[generateF01ProgressionSpec] {msg}")


def rows_after_header(ws, first_header):
    """Rows (as dicts) of the table whose header row starts with `first_header`, until a blank first cell."""
    rows = list(ws.iter_rows(values_only=True))
    for i, r in enumerate(rows):
        if r and cell(r[0]) == first_header:
            headers = [cell(h) for h in r]
            out = []
            for body in rows[i + 1:]:
                if not body or not cell(body[0]):
                    break
                out.append({h: cell(v) for h, v in zip(headers, body) if h})
            return headers, out
    fail(f"sheet {ws.title}: header row starting with {first_header!r} not found")


def main():
    src = os.path.join(ROOT, SRC_REL)
    sha256 = hashlib.sha256(open(src, "rb").read()).hexdigest()
    wb = openpyxl.load_workbook(src, data_only=True)
    missing = [s for s in REQUIRED_SHEETS if s not in wb.sheetnames]
    if missing:
        fail(f"not the D5-0 V1 workbook — missing sheets {missing} (found {wb.sheetnames})")

    # 5xSTS_ChangeEvidence — the row selected for D5.
    _, evidence = rows_after_header(wb["5xSTS_ChangeEvidence"], "evidence_id")
    selected = [e for e in evidence if e.get("selected_for_D5") == "Yes"]
    if len(selected) != 1:
        fail(f"expected exactly one selected_for_D5 evidence row, found {len(selected)}")
    ref = selected[0]
    reference_mdc_ms = round(float(ref["MDC_seconds"]) * 1000)
    change_rules = []
    for r in wb["5xSTS_ChangeEvidence"].iter_rows(values_only=True):
        if r and cell(r[0]) == "D5 runtime rule":
            change_rules.append({"field": cell(r[1]), "condition": cell(r[2]), "changeClass": cell(r[3])})
    if len(change_rules) != 3 or any(f"{reference_mdc_ms}" not in c["condition"] for c in change_rules):
        fail(f"5xSTS runtime rules do not use ±{reference_mdc_ms} ms: {change_rules}")

    # TrainingExposure
    _, exposure = rows_after_header(wb["TrainingExposure"], "state_id")
    exposure_states = [{
        "stateId": e["state_id"],
        "minDays": int(float(e["completed_days_min"])),
        "maxDays": int(float(e["completed_days_max"])),
        "code": e["system_code"],
        "labelZh": e["label_zh"],
        "plannedTrainingDays": int(float(e["planned_training_days"])),
        "status": e["status"],
    } for e in exposure]
    if any(s["status"] != "Confirmed" for s in exposure_states):
        fail("TrainingExposure rows must all be Confirmed")

    # D5_DecisionTable
    headers, decision_rows = rows_after_header(wb["D5_DecisionTable"], "rule_id")
    if "auto_decision_allowed" in headers or "auto_apply_allowed" not in headers:
        fail(f"D5_DecisionTable must use auto_apply_allowed (headers: {headers})")
    rules = []
    for r in decision_rows:
        if not r["rule_id"].startswith("R"):
            continue
        label = r["decision"]
        if label not in DECISION_CODE:
            fail(f"{r['rule_id']}: unknown / TBD decision {label!r}")
        if r["rule_status"] != "Confirmed":
            fail(f"{r['rule_id']}: rule_status is {r['rule_status']!r}, not Confirmed")
        rules.append({
            "ruleId": r["rule_id"],
            "assessmentValid": r["assessment_valid"],
            "changeClass": r["change_class"],
            "trainingState": r["training_state"],
            "limitationOrDiscomfort": r["limitation_or_discomfort"],
            "decision": DECISION_CODE[label],
            "decisionLabelZh": label,
            "autoApplyAllowed": r["auto_apply_allowed"] == "Yes",
            "exerciseAction": r["exercise_action"],
            "reason": r["recommendation_reason"],
            "evidenceType": r["evidence_type"],
            "status": r["rule_status"],
        })
    ids = [r["ruleId"] for r in rules]
    if ids != [f"R{n:03d}" for n in range(1, 12)]:
        fail(f"expected R001..R011, found {ids}")

    # ExerciseTransition — every row, status kept (the engine filters to Confirmed).
    _, transition_rows = rows_after_header(wb["ExerciseTransition"], "current_exercise_id")
    transitions = []
    for t in transition_rows:
        if t["transition_type"] not in TRANSITION_TYPES:
            fail(f"unknown transition_type {t['transition_type']!r}")
        transitions.append({
            "currentExerciseId": t["current_exercise_id"],
            "currentExerciseName": t["current_exercise_name"],
            "goalId": t["goal_id"],
            "transitionType": t["transition_type"],
            "candidateExerciseId": t["candidate_exercise_id"],
            "candidateExerciseName": t["candidate_exercise_name"],
            "evidenceType": t["evidence_type"],
            "status": t["status"],
            "sourceUrl": t.get("source_url", ""),
        })
    if not transitions:
        fail("ExerciseTransition has no rows")

    # NextCycleSelectionRules — decision-specific rows (reason templates, max_replace).
    _, selection_rows = rows_after_header(wb["NextCycleSelectionRules"], "rule_id")
    general = [s for s in selection_rows if s["rule_id"].startswith("NC-G")]
    specific = []
    rows = list(wb["NextCycleSelectionRules"].iter_rows(values_only=True))
    for i, r in enumerate(rows):
        if r and cell(r[0]) == "rule_id" and cell(r[1]) == "decision":
            hdr = [cell(h) for h in r]
            for body in rows[i + 1:]:
                if not body or not cell(body[0]).startswith("NC-"):
                    break
                d = {h: cell(v) for h, v in zip(hdr, body) if h}
                if d["decision"] not in DECISION_CODE:
                    fail(f"{d['rule_id']}: unknown decision {d['decision']!r}")
                specific.append({
                    "ruleId": d["rule_id"],
                    "decision": DECISION_CODE[d["decision"]],
                    "maxReplace": int(float(d["max_replace"])),
                    "reasonTemplate": d["reason_template"],
                    "status": d["status"],
                })
    if [s["ruleId"] for s in specific] != ["NC-M01", "NC-P01", "NC-A01", "NC-A02", "NC-N01"]:
        fail(f"unexpected decision-specific selection rules {[s['ruleId'] for s in specific]}")

    spec_source = {"file": SRC_REL, "sha256": sha256, "version": "D5-0-V1"}
    buf = io.StringIO()
    w = buf.write
    w("/**\n * GENERATED FILE — do not edit by hand.\n")
    w(f" * Source of Truth: {SRC_REL}\n * Source SHA-256:  {sha256}\n")
    w(" * Generator: scripts/generateF01ProgressionSpec.py\n *\n")
    w(" * Phase D5 (D5-0 V1) next-cycle decision rules. All rows are ReMotion\n")
    w(" * Engineering Rules / literature references — not diagnoses, clinical\n")
    w(" * cut-offs or prescriptions. The reference MDC (Yin et al. 2023, sarcopenia\n")
    w(" * subgroup) only classifies the measured 5xSTS change; it is not an MCID.\n */\n")
    w(f"export const F01_PROGRESSION_SPEC_SOURCE = Object.freeze({json.dumps(spec_source, ensure_ascii=False)});\n\n")
    w("/** 5xSTS_ChangeEvidence — selected reference. Decisions compare RAW milliseconds against referenceMdcMs. */\n")
    w("export const F01_CHANGE_REFERENCE = Object.freeze(")
    w(json.dumps({
        "evidenceId": ref["evidence_id"],
        "reference": f"{ref['reference']} ({ref['year']})",
        "group": ref["group"],
        "referenceMdcMs": reference_mdc_ms,
        "referenceMdcPercent": float(ref["MDC_percent"]),
        "referenceMdcPercentUse": "not_used_in_D5_V1",
        "populationFit": ref["population_fit"],
        "runtimeRules": change_rules,
    }, ensure_ascii=False, indent=2))
    w(");\n\n/** TrainingExposure — ReMotion product-cycle classification, not an adherence standard. */\n")
    w(f"export const F01_TRAINING_EXPOSURE_STATES = Object.freeze({json.dumps(exposure_states, ensure_ascii=False, indent=2)});\n\n")
    w("/** D5_DecisionTable R001..R011 (decision mapped to its program code). */\n")
    w(f"export const F01_D5_DECISION_RULES = Object.freeze({json.dumps(rules, ensure_ascii=False, indent=2)});\n\n")
    w("/** ExerciseTransition — ALL rows with their status; only status \"Confirmed\" may be used automatically. */\n")
    w(f"export const F01_EXERCISE_TRANSITIONS = Object.freeze({json.dumps(transitions, ensure_ascii=False, indent=2)});\n\n")
    w("/** NextCycleSelectionRules — decision-specific rows (reason templates, max replacements). */\n")
    w(f"export const F01_NEXT_CYCLE_SELECTION_RULES = Object.freeze({json.dumps(specific, ensure_ascii=False, indent=2)});\n\n")
    w("/** NextCycleSelectionRules — general rules NC-G01..NC-G11 (traceability). */\n")
    w(f"export const F01_NEXT_CYCLE_GENERAL_RULES = Object.freeze({json.dumps([{'ruleId': g['rule_id'], 'rule': g.get('rule', ''), 'status': g['status']} for g in general], ensure_ascii=False, indent=2)});\n")

    out = os.path.join(ROOT, OUT_REL)
    with open(out, "w", encoding="utf-8", newline="\n") as f:
        f.write(buf.getvalue())
    counts = {s: sum(1 for t in transitions if t["status"] == s) for s in sorted({t["status"] for t in transitions})}
    print(f"wrote {OUT_REL}: {len(rules)} decision rules, {len(transitions)} transitions {counts}, "
          f"referenceMdcMs={reference_mdc_ms}, source sha256={sha256}")


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    main()
